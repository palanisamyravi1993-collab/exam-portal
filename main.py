"""Online Examination Portal - FastAPI + SQLAlchemy (SQLite default, PostgreSQL via DATABASE_URL).
Includes Rate Limiting (100 req/hr), CORS, Safe Exam Browser (SEB), and Manual Grading.
"""
import os, io, csv, json, hmac, hashlib, base64, plistlib, gzip, time, threading
from datetime import datetime, timedelta
from fastapi import FastAPI, Depends, HTTPException, UploadFile, File, Header, Request
from fastapi.responses import StreamingResponse, Response, FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from starlette.middleware.base import BaseHTTPMiddleware
from sqlalchemy import create_engine, Column, Integer, String, Float, Boolean, DateTime, JSON, Text
from sqlalchemy.orm import declarative_base, sessionmaker, Session

# Load environment variables from .env if present
try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

SECRET = os.getenv("SECRET_KEY", "change-me-in-production-random-secret-key")
PUBLIC_URL = os.getenv("PUBLIC_URL", "http://localhost:8000").rstrip("/")
DATABASE_URL = os.getenv("DATABASE_URL", "sqlite:///./exam.db")
RATE_LIMIT_PER_HOUR = int(os.getenv("RATE_LIMIT_PER_HOUR", "100"))
CORS_ORIGINS_RAW = os.getenv("CORS_ORIGINS", "*")
CORS_ORIGINS = [o.strip() for o in CORS_ORIGINS_RAW.split(",") if o.strip()]

# Database Engine & Session
# Check if SQLite requires thread check disable
engine_kwargs = {}
if DATABASE_URL.startswith("sqlite"):
    engine_kwargs["connect_args"] = {"check_same_thread": False}

engine = create_engine(DATABASE_URL, **engine_kwargs)
SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
Base = declarative_base()
now = datetime.utcnow

# ---------- Database Models ----------
class User(Base):
    __tablename__ = "users"
    id = Column(Integer, primary_key=True)
    username = Column(String, unique=True, index=True)
    name = Column(String)
    role = Column(String)  # admin | faculty | student
    pw = Column(String)
    active = Column(Boolean, default=True)
    dept = Column(String, default="")
    klass = Column(String, default="")

class Exam(Base):
    __tablename__ = "exams"
    id = Column(Integer, primary_key=True)
    title = Column(String)
    subject = Column(String, default="")
    start = Column(DateTime)
    end = Column(DateTime)
    duration = Column(Integer)  # in minutes
    pass_pct = Column(Float, default=40)
    published = Column(Boolean, default=False)
    show_results = Column(Boolean, default=False)
    seb_required = Column(Boolean, default=False)
    seb_key = Column(String, default="")
    access_code = Column(String, default="")
    public = Column(Boolean, default=False)
    assigned = Column(JSON, default=list)
    created_by = Column(Integer)

class Question(Base):
    __tablename__ = "questions"
    id = Column(Integer, primary_key=True)
    exam_id = Column(Integer, index=True)
    pos = Column(Integer)
    text = Column(Text)
    qtype = Column(String)  # MCQ | MSQ | TF | NAT | FILL | SHORT
    options = Column(JSON, default=list)
    answer = Column(JSON, default=list)
    marks = Column(Float, default=1)
    neg = Column(Float, default=0)
    subject = Column(String, default="")

class Attempt(Base):
    __tablename__ = "attempts"
    id = Column(Integer, primary_key=True)
    exam_id = Column(Integer, index=True)
    user_id = Column(Integer, index=True)
    started = Column(DateTime, default=now)
    deadline = Column(DateTime)
    submitted = Column(DateTime, nullable=True)
    answers = Column(JSON, default=dict)
    flags = Column(JSON, default=list)
    score = Column(Float, nullable=True)
    pending = Column(Boolean, default=False)
    events = Column(JSON, default=list)

# ---------- Security & Token Handling ----------
def hash_pw(p, salt=None):
    salt = salt or os.urandom(8).hex()
    return salt + "$" + hashlib.pbkdf2_hmac("sha256", p.encode(), salt.encode(), 100_000).hex()

def check_pw(p, h):
    try:
        salt = h.split("$")[0]
        return hmac.compare_digest(hash_pw(p, salt), h)
    except Exception:
        return False

def make_token(uid):
    body = base64.urlsafe_b64encode(json.dumps({"u": uid, "e": time.time() + 8 * 3600}).encode()).decode()
    return body + "." + hmac.new(SECRET.encode(), body.encode(), hashlib.sha256).hexdigest()

def db():
    s = SessionLocal()
    try:
        yield s
    finally:
        s.close()

def me(authorization: str = Header(""), s: Session = Depends(db)):
    try:
        body, sig = authorization.removeprefix("Bearer ").split(".")
        assert hmac.compare_digest(sig, hmac.new(SECRET.encode(), body.encode(), hashlib.sha256).hexdigest())
        d = json.loads(base64.urlsafe_b64decode(body))
        assert d["e"] > time.time()
        u = s.get(User, d["u"])
        assert u and u.active
        return u
    except Exception:
        raise HTTPException(401, "Not authenticated")

def staff(u: User = Depends(me)):
    if u.role not in ("admin", "faculty"):
        raise HTTPException(403, "Staff only")
    return u

def admin(u: User = Depends(me)):
    if u.role != "admin":
        raise HTTPException(403, "Admin only")
    return u

def own_exam(eid, u, s):
    e = s.get(Exam, eid)
    if not e:
        raise HTTPException(404, "Exam not found")
    if u.role == "faculty" and e.created_by != u.id:
        raise HTTPException(403, "Not your exam")
    return e

def check_seb(e, request: Request):
    """SEB sends X-SafeExamBrowser-RequestHash = sha256(absolute URL + BrowserExamKey)."""
    if not e.seb_required:
        return
    h = request.headers.get("x-safeexambrowser-requesthash")
    if e.seb_key:
        url = PUBLIC_URL + request.url.path + ("?" + request.url.query if request.url.query else "")
        if h and hmac.compare_digest(h, hashlib.sha256((url + e.seb_key).encode()).hexdigest()):
            return
        raise HTTPException(403, "Safe Exam Browser required (invalid browser key)")
    if "SEB" not in request.headers.get("user-agent", ""):
        raise HTTPException(403, "Safe Exam Browser required")

# ---------- Automated Grading ----------
def grade(q, a):
    """Returns marks awarded, or None for manual-evaluation questions."""
    if q.qtype == "SHORT":
        return None
    if a in (None, "", []):
        return 0.0
    if q.qtype in ("MCQ", "TF"):
        return q.marks if str(a).strip().lower() == str(q.answer[0]).lower() else -q.neg
    if q.qtype == "MSQ":
        return q.marks if {str(x).lower() for x in a} == {str(x).lower() for x in q.answer} else 0.0
    for c in q.answer:  # NAT / FILL: numeric (1% tolerance) or case-insensitive text
        try:
            if abs(float(a) - float(c)) <= abs(float(c)) * 0.01 + 1e-9:
                return q.marks
        except ValueError:
            if str(a).strip().lower() == str(c).strip().lower():
                return q.marks
    return 0.0

def finalize(at, s):
    if at.submitted:
        return at
    qs = s.query(Question).filter_by(exam_id=at.exam_id).all()
    marks = [grade(q, at.answers.get(str(q.id))) for q in qs]
    at.score = round(sum(m for m in marks if m is not None), 2)
    at.pending = any(m is None for m in marks)
    at.submitted = now()
    s.commit()
    return at

def expire(at, s):
    if not at.submitted and now() > at.deadline + timedelta(seconds=5):
        finalize(at, s)

# ---------- Rate Limiter Middleware (100 req/hour) ----------
class RateLimiterMiddleware(BaseHTTPMiddleware):
    def __init__(self, app, max_requests: int = 100, window_seconds: int = 3600):
        super().__init__(app)
        self.max_requests = max_requests
        self.window_seconds = window_seconds
        self.clients = {}  # client_id -> list of timestamp floats
        self.lock = threading.Lock()

    def _get_client_id(self, request: Request) -> str:
        # Check authorization header first for authenticated client
        auth = request.headers.get("authorization", "")
        if auth and "Bearer " in auth:
            return "token:" + hashlib.sha256(auth.encode()).hexdigest()[:16]
        # Check Cloudflare IP header
        cf_ip = request.headers.get("cf-connecting-ip")
        if cf_ip:
            return "ip:" + cf_ip.strip()
        # Check standard X-Forwarded-For
        xf = request.headers.get("x-forwarded-for")
        if xf:
            return "ip:" + xf.split(",")[0].strip()
        client = request.client
        return "ip:" + (client.host if client else "unknown")

    async def dispatch(self, request: Request, call_next):
        path = request.url.path
        # Allow static files, root, and health checks to bypass rate limit
        if path.startswith("/static") or path in ("/", "/favicon.ico", "/api/health"):
            return await call_next(request)

        client_id = self._get_client_id(request)
        current_time = time.time()
        window_start = current_time - self.window_seconds

        with self.lock:
            timestamps = [t for t in self.clients.get(client_id, []) if t > window_start]
            if len(timestamps) >= self.max_requests:
                earliest = timestamps[0]
                retry_after = int(earliest + self.window_seconds - current_time) + 1
                reset_time = int(earliest + self.window_seconds)
                return JSONResponse(
                    status_code=429,
                    content={
                        "detail": f"Rate limit exceeded. Maximum {self.max_requests} requests per hour allowed.",
                        "limit": self.max_requests,
                        "retry_after_seconds": max(1, retry_after)
                    },
                    headers={
                        "Retry-After": str(max(1, retry_after)),
                        "X-RateLimit-Limit": str(self.max_requests),
                        "X-RateLimit-Remaining": "0",
                        "X-RateLimit-Reset": str(reset_time)
                    }
                )

            timestamps.append(current_time)
            self.clients[client_id] = timestamps
            remaining = self.max_requests - len(timestamps)
            reset_time = int(timestamps[0] + self.window_seconds)

        response = await call_next(request)
        response.headers["X-RateLimit-Limit"] = str(self.max_requests)
        response.headers["X-RateLimit-Remaining"] = str(max(0, remaining))
        response.headers["X-RateLimit-Reset"] = str(reset_time)
        return response

# ---------- Application Initialization ----------
app = FastAPI(title="Exam Portal", description="Modern Online Examination Portal API with 100 req/hr rate limiting and SEB support")

# Add Rate Limiting Middleware
app.add_middleware(RateLimiterMiddleware, max_requests=RATE_LIMIT_PER_HOUR, window_seconds=3600)

# Add CORS Middleware
app.add_middleware(
    CORSMiddleware,
    allow_origins=CORS_ORIGINS if CORS_ORIGINS != ["*"] else ["*"],
    allow_credentials=True if CORS_ORIGINS != ["*"] else False,
    allow_methods=["*"],
    allow_headers=["*"],
)

def total(s, eid):
    return sum(q.marks for q in s.query(Question).filter_by(exam_id=eid))

def exam_out(e, s):
    return dict(
        id=e.id, title=e.title, subject=e.subject,
        start=e.start.isoformat() + "Z", end=e.end.isoformat() + "Z",
        duration=e.duration, pass_pct=e.pass_pct, published=e.published,
        show_results=e.show_results, seb_required=e.seb_required,
        total=total(s, e.id), assigned=len(e.assigned or [])
    )

# ---------- Health Check & System Status ----------
@app.get("/api/health")
def health():
    return {
        "status": "healthy",
        "timestamp": datetime.utcnow().isoformat() + "Z",
        "rate_limit_per_hour": RATE_LIMIT_PER_HOUR,
        "database": "sqlite" if DATABASE_URL.startswith("sqlite") else "postgresql",
        "public_url": PUBLIC_URL
    }

# ---------- Auth & Account Endpoints ----------
@app.post("/api/login")
def login(d: dict, s: Session = Depends(db)):
    u = s.query(User).filter_by(username=d.get("username", "").strip()).first()
    if not u or not u.active or not check_pw(d.get("password", ""), u.pw):
        raise HTTPException(401, "Invalid credentials")
    return {"token": make_token(u.id), "user": dict(id=u.id, name=u.name, role=u.role, dept=u.dept, klass=u.klass)}

@app.post("/api/public/join")  # shared exam password -> guest student session (no pre-created account)
def join(d: dict, s: Session = Depends(db)):
    code, roll, name = ((d.get("code") or "").strip(), (d.get("roll") or "").strip(), (d.get("name") or "").strip())
    if not (code and roll and name):
        raise HTTPException(400, "Name, roll number and exam password are required")
    e = s.query(Exam).filter(Exam.public == True, Exam.published == True, Exam.access_code == code).first()
    if not e:
        raise HTTPException(403, "Wrong exam password")
    u = s.query(User).filter_by(username=roll).first()
    if u and (u.klass != "public" or not u.active):
        raise HTTPException(403, "This ID is reserved - use the normal login")
    if not u:
        u = User(username=roll, name=name, role="student", pw=hash_pw(os.urandom(8).hex()), klass="public")
        s.add(u); s.commit()
    if u.id not in (e.assigned or []):
        e.assigned = (e.assigned or []) + [u.id]; s.commit()
    return {"token": make_token(u.id), "user": dict(id=u.id, name=u.name, role=u.role, klass=u.klass)}

# ---------- User Management ----------
@app.post("/api/users")
def add_user(d: dict, a=Depends(admin), s: Session = Depends(db)):
    if d["role"] not in ("faculty", "student"):
        raise HTTPException(400, "role must be faculty or student")
    if s.query(User).filter_by(username=d["username"]).first():
        raise HTTPException(400, "Username exists")
    s.add(User(username=d["username"], name=d["name"], role=d["role"], pw=hash_pw(d["password"]), dept=d.get("dept", ""), klass=d.get("klass", "")))
    s.commit()
    return {"ok": True}

@app.post("/api/users/import")  # CSV: username,name,password,role,dept,klass
async def import_users(file: UploadFile = File(...), a=Depends(admin), s: Session = Depends(db)):
    n = 0
    for r in csv.DictReader(io.StringIO((await file.read()).decode("utf-8-sig"))):
        if r.get("username") and not s.query(User).filter_by(username=r["username"]).first():
            s.add(User(username=r["username"], name=r.get("name", ""), role=r.get("role") or "student",
                       pw=hash_pw(r.get("password") or "changeme"), dept=r.get("dept", ""), klass=r.get("klass", "")))
            n += 1
    s.commit()
    return {"imported": n}

@app.get("/api/users")
def users(a=Depends(admin), s: Session = Depends(db)):
    return [dict(id=u.id, username=u.username, name=u.name, role=u.role, active=u.active, klass=u.klass, dept=u.dept)
            for u in s.query(User).order_by(User.role, User.username)]

@app.patch("/api/users/{uid}")
def toggle_user(uid: int, d: dict, a=Depends(admin), s: Session = Depends(db)):
    u = s.get(User, uid)
    if not u: raise HTTPException(404, "User not found")
    if "active" in d: u.active = bool(d["active"])
    s.commit()
    return {"ok": True}

@app.patch("/api/users/{uid}/password")
def change_password(uid: int, d: dict, u=Depends(me), s: Session = Depends(db)):
    target = s.get(User, uid)
    if not target: raise HTTPException(404, "User not found")
    if u.role != "admin" and u.id != uid:
        raise HTTPException(403, "Not authorized to change this password")
    new_pw = d.get("password", "").strip()
    if len(new_pw) < 6:
        raise HTTPException(400, "Password must be at least 6 characters")
    target.pw = hash_pw(new_pw)
    s.commit()
    return {"ok": True}

# ---------- Exam Management ----------
@app.post("/api/exams")
def create_exam(d: dict, u=Depends(staff), s: Session = Depends(db)):
    e = Exam(
        title=d["title"], subject=d.get("subject", ""),
        start=datetime.fromisoformat(d["start"].rstrip("Z")),
        end=datetime.fromisoformat(d["end"].rstrip("Z")),
        duration=int(d["duration"]), pass_pct=float(d.get("pass_pct", 40)),
        seb_required=bool(d.get("seb_required")), seb_key=d.get("seb_key", ""),
        access_code=d.get("access_code", ""), public=bool(d.get("public")),
        assigned=[], created_by=u.id
    )
    s.add(e); s.commit()
    return exam_out(e, s)

@app.get("/api/exams")
def list_exams(u=Depends(staff), s: Session = Depends(db)):
    q = s.query(Exam) if u.role == "admin" else s.query(Exam).filter_by(created_by=u.id)
    return [{**exam_out(e, s), "access_code": e.access_code, "public": e.public} for e in q.order_by(Exam.start.desc())]

@app.patch("/api/exams/{eid}")
def edit_exam(eid: int, d: dict, u=Depends(staff), s: Session = Depends(db)):
    e = own_exam(eid, u, s)
    for k in ("published", "show_results", "seb_required", "seb_key", "access_code", "public", "duration", "pass_pct", "title", "subject"):
        if k in d: setattr(e, k, d[k])
    s.commit()
    return exam_out(e, s)

@app.post("/api/exams/{eid}/assign")
def assign(eid: int, d: dict, u=Depends(staff), s: Session = Depends(db)):
    e = own_exam(eid, u, s)
    st = s.query(User).filter_by(role="student", active=True)
    if d.get("all"):
        ids = [x.id for x in st]
    elif d.get("klass"):
        ids = [x.id for x in st.filter_by(klass=d["klass"])]
    else:
        ids = [x.id for x in st.filter(User.username.in_(d.get("usernames", [])))]
    e.assigned = sorted(set(ids))
    s.commit()
    return {"assigned": len(e.assigned)}

# ---------- Question Management ----------
HEAD = ["Question", "Type", "Option A", "Option B", "Option C", "Option D", "Correct Answer", "Marks", "Negative Marks", "Subject"]

@app.get("/api/template")
def template():
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Questions"; ws.append(HEAD)
    ws.append(["Two fair coins are tossed. P(at least one head)?", "MCQ", "1/4", "3/4", "1", "1/2", "B", 1, 0.33, "Probability"])
    ws.append(["Select the prime numbers", "MSQ", "2", "4", "5", "9", "A,C", 2, 0, "Maths"])
    ws.append(["The Earth is round.", "TF", "", "", "", "", "True", 1, 0, "General"])
    ws.append(["P(X>1) for Exp(2)? (numeric)", "NAT", "", "", "", "", "0.1353", 1, 0, "Probability"])
    ws.append(["Capital of France? (accepted answers separated by |)", "FILL", "", "", "", "", "Paris|paris", 1, 0, "General"])
    ws.append(["Explain overfitting and how regularization prevents it.", "SHORT", "", "", "", "", "", 5, 0, "AIML"])
    b = io.BytesIO(); wb.save(b); b.seek(0)
    return StreamingResponse(
        b, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=question_template.xlsx"}
    )

def infer(t, opts, ans):
    t = (t or "").upper().strip()
    if t in ("MCQ", "MSQ", "TF", "NAT", "FILL", "SHORT"): return t
    if not ans: return "SHORT"
    if not opts: return "TF" if ans.lower() in ("true", "false") else "NAT"
    return "MSQ" if "," in ans else "MCQ"

@app.post("/api/exams/{eid}/questions/import")
async def import_q(eid: int, file: UploadFile = File(...), u=Depends(staff), s: Session = Depends(db)):
    own_exam(eid, u, s); data = await file.read(); name = (file.filename or "").lower()
    if name.endswith(".csv"):
        rows = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
    else:
        import openpyxl
        it = openpyxl.load_workbook(io.BytesIO(data), read_only=True).active.iter_rows(values_only=True)
        h = [str(c or "").strip() for c in next(it)]; rows = [dict(zip(h, r)) for r in it if r and r[0]]
    pos = s.query(Question).filter_by(exam_id=eid).count(); n = 0; errs = []
    for i, r in enumerate(rows, 2):
        try:
            g = lambda k: str(r.get(k) if r.get(k) is not None else "").strip()
            opts = [g(f"Option {c}") for c in "ABCD" if g(f"Option {c}")]; ans = g("Correct Answer"); t = infer(g("Type"), opts, ans)
            if t in ("MCQ", "MSQ"):
                a = [x.strip().upper() for x in ans.split(",")]; assert a and all(x in "ABCD"[:len(opts)] for x in a), "bad answer letter"
            elif t == "TF":
                a = [ans.capitalize()]; opts = ["True", "False"]
            else:
                a = [x.strip() for x in ans.split("|") if x.strip()]
            pos += 1; n += 1
            s.add(Question(exam_id=eid, pos=pos, text=g("Question"), qtype=t, options=opts, answer=a,
                           marks=float(g("Marks") or 1), neg=float(g("Negative Marks") or 0), subject=g("Subject")))
        except Exception as ex:
            errs.append(f"Row {i}: {ex}")
    s.commit()
    return {"imported": n, "errors": errs}

@app.get("/api/exams/{eid}/questions")
def staff_questions(eid: int, u=Depends(staff), s: Session = Depends(db)):
    own_exam(eid, u, s)
    return [dict(id=q.id, text=q.text, qtype=q.qtype, options=q.options, answer=q.answer, marks=q.marks, neg=q.neg, subject=q.subject)
            for q in s.query(Question).filter_by(exam_id=eid).order_by(Question.pos)]

# ---------- Safe Exam Browser Configuration ----------
@app.get("/api/exams/{eid}/seb")
def seb_file(eid: int, u=Depends(staff), s: Session = Depends(db)):
    e = own_exam(eid, u, s)
    cfg = {
        "startURL": f"{PUBLIC_URL}/#exam={e.id}", "sendBrowserExamKey": True, "allowQuit": True,
        "quitURL": f"{PUBLIC_URL}/#done", "browserViewMode": 1, "enableAppSwitcherCheck": True,
        "allowSpellCheck": False, "allowScreenSharing": False, "allowVirtualMachine": False,
        "enableClipboard": False, "enableLogging": False, "showTaskBar": False,
        "allowBrowsingBackForward": False, "URLFilterEnable": False, "showReloadButton": False
    }
    raw = b"plnd" + gzip.compress(plistlib.dumps(cfg))
    return Response(raw, media_type="application/seb", headers={"Content-Disposition": f"attachment; filename=exam_{e.id}.seb"})

# ---------- Student Exam Flow ----------
def assigned_exam(eid, u, s):
    e = s.get(Exam, eid)
    if not e or not e.published or u.id not in (e.assigned or []):
        raise HTTPException(404, "Exam not available")
    return e

@app.get("/api/my/exams")
def my_exams(u=Depends(me), s: Session = Depends(db)):
    out = []
    for e in s.query(Exam).filter_by(published=True):
        if u.id in (e.assigned or []):
            at = s.query(Attempt).filter_by(exam_id=e.id, user_id=u.id).first()
            if at: expire(at, s)
            o = exam_out(e, s)
            o["status"] = "submitted" if at and at.submitted else ("in_progress" if at else "new")
            o["attempt"] = at.id if at else None
            if at and at.submitted and e.show_results:
                o["score"] = at.score
                o["passed"] = at.score >= (e.pass_pct / 100 * o["total"])
                o["pending"] = at.pending
            out.append(o)
    return out

@app.post("/api/exams/{eid}/start")
def start(eid: int, request: Request, u=Depends(me), s: Session = Depends(db)):
    e = assigned_exam(eid, u, s); check_seb(e, request)
    at = s.query(Attempt).filter_by(exam_id=eid, user_id=u.id).first()
    if not at:
        if not (e.start <= now() <= e.end):
            raise HTTPException(403, "Exam is not open right now")
        at = Attempt(exam_id=eid, user_id=u.id, deadline=min(now() + timedelta(minutes=e.duration), e.end))
        s.add(at); s.commit()
    expire(at, s)
    if at.submitted:
        raise HTTPException(409, "Already submitted")
    qs = [dict(id=q.id, text=q.text, qtype=q.qtype, options=q.options, marks=q.marks, neg=q.neg, subject=q.subject)
          for q in s.query(Question).filter_by(exam_id=eid).order_by(Question.pos)]
    return dict(attempt=at.id, title=e.title, subject=e.subject, duration=e.duration,
                deadline=at.deadline.isoformat() + "Z", server_now=now().isoformat() + "Z",
                questions=qs, answers=at.answers or {}, flags=at.flags or [])

def my_attempt(aid, u, s):
    at = s.get(Attempt, aid)
    if not at or at.user_id != u.id:
        raise HTTPException(404, "Attempt not found")
    expire(at, s)
    return at

@app.put("/api/attempts/{aid}")
def save(aid: int, d: dict, u=Depends(me), s: Session = Depends(db)):
    at = my_attempt(aid, u, s)
    if at.submitted:
        raise HTTPException(409, "Already submitted")
    at.answers = d.get("answers", {})
    at.flags = d.get("flags", [])
    s.commit()
    return {"ok": True}

@app.post("/api/attempts/{aid}/submit")
def submit(aid: int, d: dict, u=Depends(me), s: Session = Depends(db)):
    at = my_attempt(aid, u, s)
    if not at.submitted:
        if "answers" in d:
            at.answers = d["answers"]
        finalize(at, s)
    return {"submitted": True, "at": at.submitted.isoformat() + "Z", "score": at.score, "pending": at.pending}

@app.post("/api/attempts/{aid}/event")
def event(aid: int, d: dict, u=Depends(me), s: Session = Depends(db)):
    at = my_attempt(aid, u, s)
    event_entry = {"t": now().isoformat() + "Z", "e": str(d.get("e", ""))[:60]}
    at.events = (at.events or []) + [event_entry]
    s.commit()
    return {"ok": True}

# ---------- Results, Analytics & Evaluation ----------
def result_rows(e, s):
    names = {u.id: u for u in s.query(User)}
    tot = total(s, e.id)
    rows = []
    for at in s.query(Attempt).filter_by(exam_id=e.id):
        expire(at, s)
        u = names.get(at.user_id)
        if not u: continue
        rows.append(dict(
            attempt_id=at.id, username=u.username, name=u.name, klass=u.klass,
            status="Submitted" if at.submitted else "In progress",
            score=at.score, total=tot,
            pct=round(at.score / tot * 100, 1) if at.score is not None and tot else None,
            result=("Pass" if at.score >= e.pass_pct / 100 * tot else "Fail") if at.score is not None else "",
            pending_manual=at.pending, events=len(at.events or []),
            submitted_at=at.submitted.isoformat() + "Z" if at.submitted else None
        ))
    done = {r["username"] for r in rows}
    for uid in (e.assigned or []):
        u = names.get(uid)
        if u and u.username not in done:
            rows.append(dict(
                attempt_id=None, username=u.username, name=u.name, klass=u.klass,
                status="Not started", score=None, total=tot, pct=None, result="",
                pending_manual=False, events=0, submitted_at=None
            ))
    return sorted(rows, key=lambda r: r["username"])

@app.get("/api/exams/{eid}/results")
def results(eid: int, u=Depends(staff), s: Session = Depends(db)):
    return result_rows(own_exam(eid, u, s), s)

@app.get("/api/exams/{eid}/export")
def export(eid: int, fmt: str = "xlsx", u=Depends(staff), s: Session = Depends(db)):
    rows = result_rows(own_exam(eid, u, s), s)
    cols = ["username", "name", "klass", "status", "score", "total", "pct", "result", "pending_manual", "events", "submitted_at"]
    if fmt == "csv":
        b = io.StringIO(); w = csv.DictWriter(b, cols); w.writeheader(); w.writerows(rows)
        return Response(b.getvalue(), media_type="text/csv", headers={"Content-Disposition": f"attachment; filename=results_{eid}.csv"})
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.append(cols)
    for r in rows: ws.append([r.get(c) for c in cols])
    b = io.BytesIO(); wb.save(b); b.seek(0)
    return StreamingResponse(
        b, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=results_{eid}.xlsx"}
    )

@app.get("/api/exams/{eid}/analytics")
def analytics(eid: int, u=Depends(staff), s: Session = Depends(db)):
    own_exam(eid, u, s)
    ats = [a for a in s.query(Attempt).filter_by(exam_id=eid) if a.submitted]
    out = []
    for q in s.query(Question).filter_by(exam_id=eid).order_by(Question.pos):
        g = [grade(q, a.answers.get(str(q.id))) for a in ats]
        g = [x for x in g if x is not None]
        out.append(dict(
            q=q.pos, text=q.text[:80], qtype=q.qtype, marks=q.marks,
            attempted=sum(1 for a in ats if a.answers.get(str(q.id)) not in (None, "", [])),
            correct=sum(1 for x in g if x == q.marks),
            pct_correct=round(100 * sum(1 for x in g if x == q.marks) / len(g), 1) if g else None
        ))
    return out

@app.get("/api/exams/{eid}/attempts/{aid}")
def attempt_detail(eid: int, aid: int, u=Depends(staff), s: Session = Depends(db)):
    own_exam(eid, u, s)
    at = s.get(Attempt, aid)
    if not at or at.exam_id != eid:
        raise HTTPException(404, "Attempt not found")
    student = s.get(User, at.user_id)
    qs = s.query(Question).filter_by(exam_id=eid).order_by(Question.pos).all()
    q_data = []
    for q in qs:
        ans = at.answers.get(str(q.id))
        g = grade(q, ans)
        q_data.append({
            "id": q.id, "pos": q.pos, "text": q.text, "qtype": q.qtype,
            "options": q.options, "correct_answer": q.answer,
            "student_answer": ans, "marks": q.marks, "neg": q.neg,
            "awarded": g
        })
    return {
        "attempt_id": at.id,
        "student": {"id": student.id, "name": student.name, "username": student.username, "klass": student.klass},
        "score": at.score,
        "submitted": at.submitted.isoformat() + "Z" if at.submitted else None,
        "pending": at.pending,
        "events": at.events or [],
        "questions": q_data
    }

@app.post("/api/exams/{eid}/attempts/{aid}/grade-manual")
def grade_manual(eid: int, aid: int, d: dict, u=Depends(staff), s: Session = Depends(db)):
    own_exam(eid, u, s)
    at = s.get(Attempt, aid)
    if not at or at.exam_id != eid:
        raise HTTPException(404, "Attempt not found")
    marks_dict = d.get("marks", {})  # {"question_id": score}
    qs = s.query(Question).filter_by(exam_id=eid).all()
    total_score = 0.0
    has_pending = False
    for q in qs:
        qid_str = str(q.id)
        if qid_str in marks_dict:
            try:
                total_score += float(marks_dict[qid_str])
            except (ValueError, TypeError):
                pass
        else:
            g = grade(q, at.answers.get(qid_str))
            if g is None:
                has_pending = True
            else:
                total_score += g
    at.score = round(total_score, 2)
    at.pending = has_pending
    s.commit()
    return {"ok": True, "score": at.score, "pending": at.pending}

# ---------- Startup Database Seeding ----------
from seed import seed

@app.on_event("startup")
def startup():
    Base.metadata.create_all(engine)
    s = SessionLocal()
    if not s.query(User).count():
        seed(s, hash_pw, User, Exam, Question)
    s.close()

# ---------- Static File Serving ----------
app.mount("/static", StaticFiles(directory="static"), name="static")

@app.get("/")
def index():
    return FileResponse("static/index.html")
